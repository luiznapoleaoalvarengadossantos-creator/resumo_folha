from datetime import datetime
import json
import os
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
import streamlit as st

ARQUIVO_DADOS = "sistema_escala.json"

# Configuração da página
st.set_page_config(page_title="Sistema de Escala", page_icon="📊", layout="wide")

# Funções de Dados
def carregar_dados():
  if os.path.exists(ARQUIVO_DADOS):
    with open(ARQUIVO_DADOS, "r", encoding="utf-8") as f:
      return json.load(f)
  return {
      "usuarios": {
          "admin": {
              "senha": "123",
              "tipo": "master",
              "nome": "Administrador Master",
              "setor": "Geral",
          }
      },
      "colaboradores": [],
      "registros": [],
  }

def salvar_dados(dados):
  with open(ARQUIVO_DADOS, "w", encoding="utf-8") as f:
    json.dump(dados, f, ensure_ascii=False, indent=4)

# Inicializar Estado da Sessão
if "dados" not in st.session_state:
  st.session_state.dados = carregar_dados()

if "usuario_logado" not in st.session_state:
  st.session_state.usuario_logado = None

if "tipo_usuario" not in st.session_state:
  st.session_state.tipo_usuario = None

dados = st.session_state.dados

# --- TELA DE LOGIN ---
if not st.session_state.usuario_logado:
  st.title("🔐 Sistema de Escala - Login")
  
  with st.form("form_login"):
    usuario_input = st.text_input("Usuário").strip().lower()
    senha_input = st.text_input("Senha", type="password").strip()
    btn_login = st.form_submit_button("Entrar")

    if btn_login:
      if usuario_input in dados["usuarios"] and dados["usuarios"][usuario_input]["senha"] == senha_input:
        st.session_state.usuario_logado = usuario_input
        st.session_state.tipo_usuario = dados["usuarios"][usuario_input]["tipo"]
        st.success(f"Bem-vindo(a), {dados['usuarios'][usuario_input]['nome']}!")
        st.rerun()
      else:
        st.error("❌ Usuário ou senha incorretos!")

else:
  user_info = dados["usuarios"][st.session_state.usuario_logado]
  
  # Barra lateral para controle de conta
  st.sidebar.title(f"👤 {user_info['nome']}")
  st.sidebar.write(f"**Tipo:** {user_info['tipo'].capitalize()}")
  st.sidebar.write(f"**Setor:** {user_info['setor']}")
  
  if st.sidebar.button("🚪 Sair / Trocar de Conta"):
    st.session_state.usuario_logado = None
    st.session_state.tipo_usuario = None
    st.rerun()

  st.sidebar.markdown("---")

  # --- PAINEL DO ADMINISTRADOR MASTER ---
  if st.session_state.tipo_usuario == "master":
    st.title("👑 Painel do Administrador Master")
    
    menu_master = st.selectbox(
        "Escolha uma Ação",
        [
            "Visualizar Escala Geral",
            "Cadastrar Novo Coordenador",
            "Cadastrar Novo Servidor",
            "Alterar Setor de Servidor",
            "Importar Colaboradores via JSON",
            "Gerar Relatórios"
        ]
    )

    if menu_master == "Visualizar Escala Geral":
      st.subheader("📊 Relatório Geral da Escala")
      if not dados["registros"]:
        st.info("Nenhum registro de escala encontrado.")
      else:
        for reg in dados["registros"]:
          with st.expander(f"Colaborador: {reg['nome']} | Setor: {reg['setor']}"):
            st.write(f"**Coordenador Responsável:** {reg['coordenador_nome']}")
            st.write(f"**HE 50%:** {reg['he_50']}h | **HE 100%:** {reg['he_100']}h (Just.: {reg['justificativa_he'] or 'N/A'})")
            st.write(f"**Sobreaviso Líquido:** {reg['sobreaviso_liquido']}h (Úteis: {reg['dias_uteis']}d, Não Úteis: {reg['dias_nao_uteis']}d)")
            st.write(f"**Adc. Noturno:** {reg['adicional_noturno']} | **Falta:** {reg['falta']} | **Férias:** {reg['ferias']} | **Atestado:** {reg['atestado']}")
            if reg.get("historico_alteracao"):
              st.caption(f"⚠️ Status: {reg['historico_alteracao']}")

    elif menu_master == "Cadastrar Novo Coordenador":
      st.subheader("➕ Cadastrar Novo Coordenador")
      with st.form("form_cad_coord"):
        novo_user = st.text_input("Nome de usuário (Login)").strip().lower()
        novo_nome = st.text_input("Nome Completo").strip()
        novo_setor = st.text_input("Setor do Coordenador").strip()
        nova_senha = st.text_input("Senha Inicial", type="password").strip()
        cadastrar_btn = st.form_submit_button("Salvar Coordenador")

        if cadastrar_btn:
          if not novo_user or not novo_nome or not novo_setor or not nova_senha:
            st.warning("Preencha todos os campos.")
          elif novo_user in dados["usuarios"]:
            st.error(f"O usuário '{novo_user}' já existe.")
          else:
            dados["usuarios"][novo_user] = {
                "senha": nova_senha,
                "tipo": "coordenador",
                "nome": novo_nome,
                "setor": novo_setor,
            }
            salvar_dados(dados)
            st.success(f"Coordenador '{novo_nome}' cadastrado com sucesso!")

    elif menu_master == "Cadastrar Novo Servidor":
      st.subheader("👤 Cadastrar Novo Servidor")
      with st.form("form_cad_servidor"):
        s_nome = st.text_input("Nome Completo do Servidor").strip().upper()
        s_setor = st.text_input("Setor / Lotação").strip()
        s_btn = st.form_submit_button("Cadastrar Servidor")

        if s_btn:
          if not s_nome or not s_setor:
            st.warning("Preencha todos os campos.")
          elif any(c["nome"].lower() == s_nome.lower() for c in dados["colaboradores"]):
            st.error(f"O servidor '{s_nome}' já está cadastrado.")
          else:
            dados["colaboradores"].append({"nome": s_nome, "setor": s_setor})
            salvar_dados(dados)
            st.success(f"Servidor '{s_nome}' cadastrado com sucesso!")

    elif menu_master == "Alterar Setor de Servidor":
      st.subheader("🔄 Alterar Setor de um Servidor")
      if not dados["colaboradores"]:
        st.info("Nenhum colaborador cadastrado.")
      else:
        nomes_colabs = [c["nome"] for c in dados["colaboradores"]]
        colab_selecionado = st.selectbox("Selecione o Servidor", nomes_colabs)
        
        servidor_obj = next(c for c in dados["colaboradores"] if c["nome"] == colab_selecionado)
        st.write(f"Setor Atual: **{servidor_obj['setor']}**")
        
        novo_setor_serv = st.text_input("Novo Setor").strip()
        if st.button("Atualizar Setor"):
          if not novo_setor_serv:
            st.warning("Informe o novo setor.")
          else:
            servidor_obj["setor"] = novo_setor_serv
            for reg in dados["registros"]:
              if reg["nome"].lower() == servidor_obj["nome"].lower():
                reg["setor"] = novo_setor_serv
            salvar_dados(dados)
            st.success("Setor alterado com sucesso!")
            st.rerun()

    elif menu_master == "Importar Colaboradores via JSON":
      st.subheader("📂 Importar Colaboradores via JSON")
      caminho_json = st.text_input("Caminho ou Nome do Arquivo JSON (ex: funcionarios.json)").strip()
      if st.button("Importar"):
        if os.path.exists(caminho_json):
          try:
            with open(caminho_json, "r", encoding="utf-8") as f:
              lista_imp = json.load(f)
            contador = 0
            for item in lista_imp:
              n = str(item.get("nome", "")).strip()
              s = str(item.get("setor", "")).strip() or str(item.get("lotacao", "")).strip()
              if n and not any(c["nome"].lower() == n.lower() for c in dados["colaboradores"]):
                dados["colaboradores"].append({"nome": n, "setor": s})
                contador += 1
            salvar_dados(dados)
            st.success(f"{contador} novos colaboradores importados com sucesso!")
          except Exception as e:
            st.error(f"Erro ao ler JSON: {e}")
        else:
          st.error("Arquivo não encontrado.")

    elif menu_master == "Gerar Relatórios":
      st.subheader("📊 Central de Relatórios")
      if not dados["registros"]:
        st.info("Não há registros suficientes para gerar relatórios.")
      else:
        agora = datetime.now()
        timestamp_arquivo = agora.strftime("%d-%m-%Y_%H-%M-%S")
        
        col1, col2, col3 = st.columns(3)
        
        with col1:
          if st.button("Gerar CSV Geral"):
            nome_csv = f"relatorio_escala_{timestamp_arquivo}.csv"
            with open(nome_csv, mode="w", newline="", encoding="utf-8-sig") as arquivo:
              escritor = csv.writer(arquivo, delimiter=";")
              escritor.writerow(["Colaborador", "Setor", "Coordenador", "HE 50%", "HE 100%", "Sobreaviso Líq (h)"])
              for reg in dados["registros"]:
                escritor.writerow([reg["nome"], reg["setor"], reg["coordenador_nome"], reg["he_50"], reg["he_100"], reg["sobreaviso_liquido"]])
            st.success(f"Arquivo CSV gerado: {nome_csv}")

        with col2:
          if st.button("Gerar XLSX Geral"):
            nome_xlsx = f"relatorio_escala_{timestamp_arquivo}.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["Colaborador", "Setor", "Coordenador", "HE 50%", "HE 100%", "Sobreaviso Líq (h)"])
            for reg in dados["registros"]:
              ws.append([reg["nome"], reg["setor"], reg["coordenador_nome"], reg["he_50"], reg["he_100"], reg["sobreaviso_liquido"]])
            wb.save(nome_xlsx)
            st.success(f"Arquivo Excel gerado: {nome_xlsx}")

  # --- PAINEL DO COORDENADOR ---
  elif st.session_state.tipo_usuario == "coordenador":
    setor_coord = user_info["setor"]
    nome_coord = user_info["nome"]
    
    st.title(f"📝 Painel do Coordenador: {nome_coord}")
    st.markdown(f"**Setor Gerenciado:** `{setor_coord}`")

    menu_coord = st.selectbox(
        "Ações do Setor",
        ["Registrar / Alterar Ocorrências", "Visualizar Escala do Setor", "Gerar Relatórios do Setor"]
    )

    colab_setor = [c for c in dados["colaboradores"] if c["setor"].lower() == setor_coord.lower()]

    if menu_coord == "Registrar / Alterar Ocorrências":
      st.subheader("📝 Lançamento de Ocorrências")
      if not colab_setor:
        st.warning(f"Não há colaboradores cadastrados no setor '{setor_coord}'. Verifique com o Administrador Master se o setor está idêntico.")
      else:
        nomes_setor = [c["nome"] for c in colab_setor]
        colab_escolhido = st.selectbox("Selecione o Colaborador", nomes_setor)

        # Verificar se já existe registro
        reg_existente = next((r for r in dados["registros"] if r["nome"] == colab_escolhido and r["setor"] == setor_coord), None)

        with st.form("form_ocorrencia"):
          he_50_val = st.number_input("Horas Extras 50% (h)", min_value=0.0, step=0.5, value=float(reg_existente["he_50"]) if reg_existente else 0.0)
          he_100_val = st.number_input("Horas Extras 100% (h)", min_value=0.0, step=0.5, value=float(reg_existente["he_100"]) if reg_existente else 0.0)
          
          total_he_form = he_50_val + he_100_val
          just_he_val = st.text_input("Justificativa para as Horas Extras", value=reg_existente["justificativa_he"] if reg_existente else "")

          dias_u = st.number_input("Dias de Sobreaviso em Dias ÚTEIS (x 12h)", min_value=0, step=1, value=int(reg_existente["dias_uteis"]) if reg_existente else 0)
          dias_nu = st.number_input("Dias de Sobreaviso em Dias NÃO ÚTEIS (x 24h)", min_value=0, step=1, value=int(reg_existente["dias_nao_uteis"]) if reg_existente else 0)

          ad_not = st.selectbox("Adicional Noturno", ["Não", "Sim", "Detalhes"], index=0 if not reg_existente else 0)
          falta_s = st.selectbox("Possui Falta?", ["Não", "Sim"], index=0)
          ferias_s = st.selectbox("Está de Férias?", ["Não", "Sim"], index=0)
          atestado_s = st.text_input("Atestado Médico (Sim/Não/Dias)", value=reg_existente["atestado"] if reg_existente else "Não")

          btn_salvar_oc = st.form_submit_button("Salvar / Atualizar Ocorrência")

          if btn_salvar_oc:
            sobreaviso_bruto = (dias_u * 12) + (dias_nu * 24)
            sobreaviso_liquido = max(0.0, sobreaviso_bruto - total_he_form)

            novo_reg = {
                "coordenador": st.session_state.usuario_logado,
                "coordenador_nome": nome_coord,
                "nome": colab_escolhido,
                "setor": setor_coord,
                "he_50": he_50_val,
                "he_100": he_100_val,
                "justificativa_he": just_he_val,
                "dias_uteis": dias_u,
                "dias_nao_uteis": dias_nu,
                "sobreaviso_liquido": sobreaviso_liquido,
                "adicional_noturno": ad_not,
                "falta": falta_s,
                "ferias": ferias_s,
                "atestado": atestado_s,
                "historico_alteracao": "Atualizado/Registrado por Coordenador" if reg_existente else "Registro Original"
            }

            # Remover antigo se existir e adicionar o atualizado
            dados["registros"] = [r for r in dados["registros"] if not (r["nome"] == colab_escolhido and r["setor"] == setor_coord)]
            dados["registros"].append(novo_reg)
            salvar_dados(dados)
            st.success(f"Ocorrência de {colab_escolhido} salva com sucesso! (Sobreaviso Líq.: {sobreaviso_liquido}h)")

    elif menu_coord == "Visualizar Escala do Setor":
      st.subheader(f"📊 Escala Registrada - {setor_coord}")
      registros_setor = [r for r in dados["registros"] if r["setor"].lower() == setor_coord.lower()]
      if not registros_setor:
        st.info("Nenhum registro encontrado para o seu setor.")
      else:
        for reg in registros_setor:
          with st.container():
            st.write(f"**Servidor:** {reg['nome']}")
            st.write(f"• HE 50%: {reg['he_50']}h | HE 100%: {reg['he_100']}h")
            st.write(f"• Sobreaviso Líquido: {reg['sobreaviso_liquido']}h (Úteis: {reg['dias_uteis']}d, Não Úteis: {reg['dias_nao_uteis']}d)")
            st.write(f"• Adc. Noturno: {reg['adicional_noturno']} | Falta: {reg['falta']} | Férias: {reg['ferias']} | Atestado: {reg['atestado']}")
            st.markdown("---")

    elif menu_coord == "Gerar Relatórios do Setor":
      st.subheader("📊 Exportar Relatórios")
      registros_setor = [r for r in dados["registros"] if r["setor"].lower() == setor_coord.lower()]
      if not registros_setor:
        st.info("Não há dados suficientes no seu setor para exportar.")
      else:
        setor_limpo = setor_coord.replace(" ", "_").replace(".", "")
        timestamp_arquivo = datetime.now().strftime("%d-%m-%Y_%H-%M-%S")
        
        if st.button("Gerar Relatório XLSX do Setor"):
          nome_xlsx = f"relatorio_setor_{setor_limpo}_{timestamp_arquivo}.xlsx"
          wb = openpyxl.Workbook()
          ws = wb.active
          ws.append(["Colaborador", "Setor", "HE 50%", "HE 100%", "Dias Úteis", "Dias Não Úteis", "Sobreaviso Líq (h)"])
          for reg in registros_setor:
            ws.append([reg["nome"], reg["setor"], reg["he_50"], reg["he_100"], reg["dias_uteis"], reg["dias_nao_uteis"], reg["sobreaviso_liquido"]])
          wb.save(nome_xlsx)
          st.success(f"Relatório Excel gerado com sucesso: {nome_xlsx}")